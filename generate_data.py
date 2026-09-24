"""
Genera un Excel de ventas FALSO para probar el generador de reportes.

Los datos son inventados: no hay personas reales.
Lleva "suciedad" a propósito (filas vacías, duplicados, nombres mal escritos)
para demostrar que generate_report.py sabe limpiarla.

Uso:
    python generate_data.py
    python generate_data.py --output data/otro_archivo.xlsx
"""

import argparse
import random
import re
import zipfile
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# Semilla fija: el azar sale siempre igual, así el Excel es reproducible
SEED = 42

# Periodo de ventas: enero a junio de 2026
START_DATE = date(2026, 1, 1)
END_DATE = date(2026, 6, 30)

# Fecha fija que se escribe "por dentro" del .xlsx (ver make_reproducible)
FIXED_TIMESTAMP = "2026-07-01T00:00:00Z"
FIXED_ZIP_TIME = (2026, 7, 1, 0, 0, 0)

# Cuántas líneas de pedido limpias queremos (luego se suma la suciedad)
TARGET_ROWS = 1180

# Catálogo: producto -> (categoría, precio unitario en euros)
PRODUCTS = {
    "Portátil 14 pulgadas": ("Informática", 649.00),
    "Monitor 27 pulgadas": ("Informática", 229.90),
    "Teclado mecánico": ("Informática", 79.95),
    "Ratón inalámbrico": ("Informática", 24.99),
    "Disco SSD 1TB": ("Informática", 89.50),
    "Silla ergonómica": ("Oficina", 189.00),
    "Mesa elevable": ("Oficina", 349.00),
    "Lámpara LED de escritorio": ("Oficina", 34.90),
    "Pack 500 folios": ("Oficina", 5.49),
    "Auriculares Bluetooth": ("Audio", 59.99),
    "Altavoz portátil": ("Audio", 45.00),
    "Micrófono USB": ("Audio", 69.90),
    "Cafetera espresso": ("Hogar", 119.00),
    "Hervidor eléctrico": ("Hogar", 29.95),
    "Purificador de aire": ("Hogar", 159.00),
}

# Peso de cada producto: los baratos se venden más a menudo
PRODUCT_WEIGHTS = [3, 4, 6, 10, 6, 3, 2, 7, 12, 8, 6, 4, 4, 6, 3]

CITIES = ["Madrid", "Barcelona", "Valencia", "Sevilla", "Bilbao", "Zaragoza", "Málaga"]
CITY_WEIGHTS = [30, 25, 12, 10, 8, 8, 7]

# Efecto temporada: multiplica cuántos pedidos hay cada mes (1 = normal)
MONTH_FACTOR = {1: 1.2, 2: 0.8, 3: 0.9, 4: 1.0, 5: 1.1, 6: 1.3}

# Cantidad de suciedad a meter
N_EMPTY_ROWS = 15
N_DUPLICATES = 12
MESSY_NAME_RATE = 0.08  # 8 % de filas con el nombre del producto mal escrito


def random_order_date() -> date:
    """Elige un día al azar, con más probabilidad en los meses fuertes."""
    total_days = (END_DATE - START_DATE).days + 1
    while True:
        day = START_DATE + timedelta(days=random.randrange(total_days))
        # Truco de "aceptar o rechazar": cuanto más alto el factor del mes,
        # más fácil es que el día se quede. Así cada mes tiene su peso.
        if random.random() < MONTH_FACTOR[day.month] / max(MONTH_FACTOR.values()):
            return day


def mess_up_name(name: str) -> str:
    """Estropea un nombre como lo haría una persona tecleando a mano."""
    variants = [
        f"  {name}",                  # espacios delante
        f"{name}   ",                 # espacios detrás
        name.upper(),                 # TODO EN MAYÚSCULAS
        name.lower(),                 # todo en minúsculas
        name.replace(" ", "  ", 1),   # doble espacio en medio
    ]
    return random.choice(variants)


def build_clean_rows() -> list[dict]:
    """Crea pedidos correctos. Un pedido puede tener varias líneas (productos)."""
    rows = []
    order_number = 1
    product_names = list(PRODUCTS)

    while len(rows) < TARGET_ROWS:
        order_id = f"PED-{order_number:05d}"  # id provisional, se renumera al final
        order_date = random_order_date()
        city = random.choices(CITIES, weights=CITY_WEIGHTS)[0]

        # Cada pedido lleva de 1 a 4 productos distintos (casi siempre pocos)
        n_lines = random.choices([1, 2, 3, 4], weights=[50, 30, 15, 5])[0]
        chosen = set()
        while len(chosen) < n_lines:
            chosen.add(random.choices(product_names, weights=PRODUCT_WEIGHTS)[0])

        for product in sorted(chosen):
            category, price = PRODUCTS[product]
            rows.append({
                "fecha": order_date,
                "pedido_id": order_id,
                "producto": product,
                "categoria": category,
                "cantidad": random.choices([1, 2, 3, 5, 10], weights=[60, 20, 10, 7, 3])[0],
                "precio_unitario": price,
                "ciudad": city,
            })
        order_number += 1

    # Ordenamos por fecha: así parece un registro real, día a día
    rows.sort(key=lambda r: (r["fecha"], r["pedido_id"]))

    # Renumeramos los pedidos en orden de fecha (el primero de enero es PED-00001),
    # como haría un programa de ventas real
    new_ids = {}
    for row in rows:
        if row["pedido_id"] not in new_ids:
            new_ids[row["pedido_id"]] = f"PED-{len(new_ids) + 1:05d}"
        row["pedido_id"] = new_ids[row["pedido_id"]]
    return rows


def add_dirt(rows: list[dict]) -> list[dict]:
    """Mete suciedad realista: nombres mal escritos, duplicados y filas vacías."""
    dirty = [dict(r) for r in rows]  # copia, para no tocar la lista original

    # 1) Nombres de producto mal escritos
    for row in dirty:
        if random.random() < MESSY_NAME_RATE:
            row["producto"] = mess_up_name(row["producto"])

    # 2) Duplicados: la misma fila copiada justo debajo (típico copiar-pegar)
    for _ in range(N_DUPLICATES):
        pos = random.randrange(len(dirty))
        dirty.insert(pos + 1, dict(dirty[pos]))

    # 3) Filas vacías sueltas por el medio
    empty_row = {col: None for col in dirty[0]}
    for _ in range(N_EMPTY_ROWS):
        pos = random.randrange(1, len(dirty))
        dirty.insert(pos, dict(empty_row))

    return dirty


def save_excel(rows: list[dict], output_path: Path) -> None:
    """Guarda las filas en Excel con columnas legibles."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(rows)
    df.to_excel(output_path, index=False, sheet_name="ventas")

    # Retoque visual con openpyxl: ancho de columnas y formato de fecha/precio.
    # No cambia los datos, solo cómo se ven al abrir el Excel.
    wb = load_workbook(output_path)
    ws = wb["ventas"]
    widths = {"A": 12, "B": 12, "C": 30, "D": 14, "E": 10, "F": 16, "G": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"   # fecha
        row[5].number_format = "#,##0.00"     # precio_unitario
    ws.freeze_panes = "A2"  # la cabecera se queda fija al bajar
    wb.save(output_path)
    make_reproducible(output_path)


def make_reproducible(xlsx_path: Path) -> None:
    """Quita la hora actual de dentro del .xlsx para que salga igual byte a byte.

    Un .xlsx es un ZIP con archivos XML. La hora se cuela en dos sitios:
    - docProps/core.xml guarda "creado" y "modificado" = ahora
    - cada archivo del ZIP lleva su propia hora de guardado
    Si no lo arreglamos, git ve el Excel "cambiado" cada vez que lo regeneramos,
    aunque los datos sean idénticos.
    """
    with zipfile.ZipFile(xlsx_path) as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]

    with zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            if name == "docProps/core.xml":
                # Cambiamos el contenido de <dcterms:created> y <dcterms:modified>
                text = data.decode("utf-8")
                text = re.sub(
                    r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)",
                    rf"\g<1>{FIXED_TIMESTAMP}\g<2>",
                    text,
                )
                data = text.encode("utf-8")
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera un Excel de ventas de ejemplo (datos falsos).")
    parser.add_argument(
        "--output",
        default="data/sample_sales.xlsx",
        help="Ruta del Excel a crear (por defecto: data/sample_sales.xlsx)",
    )
    args = parser.parse_args()

    random.seed(SEED)
    clean_rows = build_clean_rows()
    dirty_rows = add_dirt(clean_rows)
    save_excel(dirty_rows, Path(args.output))

    n_orders = len({r["pedido_id"] for r in clean_rows})
    print(f"Excel creado: {args.output}")
    print(f"  Filas totales: {len(dirty_rows)}")
    print(f"  Líneas limpias: {len(clean_rows)} ({n_orders} pedidos)")
    print(f"  Suciedad: {N_DUPLICATES} duplicados, {N_EMPTY_ROWS} filas vacías, "
          f"~{MESSY_NAME_RATE:.0%} nombres mal escritos")


if __name__ == "__main__":
    main()
